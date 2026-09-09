"""Generate All_TMPs.xlsx style hourly performance workbook."""

from __future__ import annotations

import os
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from app.reports.catalog import CPU_FORMAT_COLUMNS, IOPS_FORMAT_COLUMNS, NUMBER_FORMAT_COLUMNS

try:
    from PIL import Image as PILImage  # type: ignore
except Exception:  # pragma: no cover
    PILImage = None

try:
    from openpyxl.drawing.image import Image as XLImage
except Exception:  # pragma: no cover
    XLImage = None

TMP_COLUMNS = [
    "DateTime",
    "Latency",
    "Read Latency",
    "Write Latency",
    "Avg. Size",
    "Read Size",
    "Write Size",
    "Total IOPS",
    "Read IOPS",
    "Write IOPS",
    "CPU Utilization",
]

TMP_RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def _clean_location(location: str) -> str:
    text = location.replace(" PowerStore", "").replace(" PoweStore", "")
    normalized = unicodedata.normalize("NFKD", text)
    return normalized.encode("ascii", "ignore").decode("ascii")


def _format_cpu(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).replace("%", "").strip()
    if not text:
        return ""
    try:
        return f"%{float(text):.1f}"
    except (TypeError, ValueError):
        return ""


def _format_iops(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    try:
        return f"{float(value) / 1000:.3f}"
    except (TypeError, ValueError):
        return ""


def _format_number(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    try:
        number = float(value)
        if number.is_integer():
            return int(number)
        return number
    except (TypeError, ValueError):
        return None


def build_server_tmp_table(
    location: str,
    server: str,
    df: pd.DataFrame,
    columns: list[str] | None = None,
) -> pd.DataFrame:
    """Build one server TMP table matching All_TMPs.xlsx layout."""
    location_clean = _clean_location(location)
    work = df.copy()
    if "Timestamp" in work.columns:
        work = work.rename(columns={"Timestamp": "DateTime"})
    work["DateTime"] = work["DateTime"].astype(str)

    selected = columns or list(TMP_COLUMNS)
    if not selected or selected[0] != "DateTime":
        selected = ["DateTime"] + [column for column in selected if column != "DateTime"]
    for column in selected:
        if column not in work.columns:
            work[column] = None
    work = work[selected]

    formatted = work.copy()
    for column in selected:
        if column in NUMBER_FORMAT_COLUMNS:
            formatted[column] = formatted[column].apply(_format_number)
        elif column in IOPS_FORMAT_COLUMNS:
            formatted[column] = formatted[column].apply(_format_iops)
        elif column in CPU_FORMAT_COLUMNS:
            formatted[column] = formatted[column].apply(_format_cpu)

    header = [f"{location_clean} {server}"] + selected[1:]
    header_df = pd.DataFrame([header], columns=selected)
    return pd.concat([header_df, formatted], ignore_index=True)


def _style_tmp_sheet(worksheet) -> None:
    for idx, row in enumerate(worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1), 1):
        fill = TMP_RED_FILL if idx % 2 == 1 else WHITE_FILL
        for cell in row:
            cell.fill = fill
            cell.border = THIN_BORDER

    for col_idx, cell in enumerate(worksheet[1], 1):
        if cell.value and "CPU Utilization" in str(cell.value):
            for row in worksheet.iter_rows(min_row=1, min_col=col_idx, max_col=col_idx):
                for item in row:
                    item.font = Font(bold=True)

    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value is not None and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)


def _add_cover_sheet(workbook: Workbook, static_dir: Path | None) -> None:
    cover = workbook.create_sheet(title="Cover", index=0)
    cover.merge_cells("C10:M13")
    cover["C10"] = "Vodafone All Locations Storage Performance Report"
    cover["C10"].font = Font(size=28, bold=True)
    cover["C10"].alignment = Alignment(horizontal="center", vertical="center")

    for row in cover.iter_rows():
        for cell in row:
            cell.border = None

    cover.row_dimensions[10].height = 40
    for col in range(3, 14):
        cover.column_dimensions[chr(64 + col)].width = 18

    if static_dir is None or XLImage is None:
        return

    target_logo_height = 240
    for filename, anchor in (("vodafone_logo.png", "H2"), ("ngtech_logo.png", "H18")):
        logo_path = static_dir / filename
        if not logo_path.is_file():
            continue
        image = XLImage(str(logo_path))
        if PILImage is not None:
            with PILImage.open(logo_path) as pil_image:
                orig_w, orig_h = pil_image.size
                target_w = int(orig_w * (target_logo_height / orig_h))
                image.width = target_w
                image.height = target_logo_height
        cover.add_image(image, anchor)


class HourlyReportGenerator:
    """Write hourly server TMP sheets into All_TMPs.xlsx."""

    def __init__(
        self,
        output_dir: str | os.PathLike[str],
        location_servers: dict[str, list[str]],
        server_data: dict[str, pd.DataFrame],
        static_dir: Path | None = None,
        columns: list[str] | None = None,
        filename: str = "All_TMPs.xlsx",
    ) -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.location_servers = location_servers
        self.server_data = server_data
        self.static_dir = static_dir
        self.columns = columns or list(TMP_COLUMNS)
        self.filename = filename

    def generate(self) -> str:
        workbook = Workbook()
        workbook.remove(workbook.active)
        _add_cover_sheet(workbook, self.static_dir)

        sheets_written = 0
        for location, servers in self.location_servers.items():
            location_clean = _clean_location(location)
            for server in servers:
                df = self.server_data.get(server)
                if df is None or df.empty:
                    continue
                tmp_df = build_server_tmp_table(location, server, df, columns=self.columns)
                sheet_name = f"{location_clean}_{server}"[:31]
                worksheet = workbook.create_sheet(title=sheet_name)
                for row in dataframe_to_rows(tmp_df, index=False, header=False):
                    worksheet.append(row)
                _style_tmp_sheet(worksheet)
                sheets_written += 1

        if sheets_written == 0:
            raise ValueError("No hourly server data available for TMP report")

        output_file = self.output_dir / self.filename
        workbook.save(output_file)
        return str(output_file)
