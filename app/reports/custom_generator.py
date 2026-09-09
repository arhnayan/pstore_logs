"""Generate custom Vodafone-style Excel reports from live hourly metrics."""

from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.workbook import Workbook

from app.reports.catalog import (
    combined_tmp_columns,
    metric_by_id,
    resolve_metric_columns,
    resolve_psk_columns,
)
from app.reports.hourly_generator import (
    HourlyReportGenerator,
    TMP_COLUMNS,
    _clean_location,
    _style_tmp_sheet,
    build_server_tmp_table,
)

HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
BOLD = Font(bold=True)


def _sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "", name).strip() or "Sheet"
    return cleaned[:31]


def _as_datetime_frame(df: pd.DataFrame) -> pd.DataFrame:
    work = df.copy()
    if "Timestamp" in work.columns and "DateTime" not in work.columns:
        work = work.rename(columns={"Timestamp": "DateTime"})
    if "DateTime" in work.columns:
        work["DateTime"] = work["DateTime"].astype(str)
    return work


def _auto_width(worksheet) -> None:
    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value is not None and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)


def _style_metric_sheet(worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = BOLD
        cell.border = THIN_BORDER
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
    for col_idx, cell in enumerate(worksheet[1], 1):
        if cell.value and "CPU Utilization" in str(cell.value):
            for row in worksheet.iter_rows(min_row=1, min_col=col_idx, max_col=col_idx):
                for item in row:
                    item.font = Font(bold=True)
    _auto_width(worksheet)


def _select_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    work = _as_datetime_frame(df)
    for column in columns:
        if column not in work.columns:
            work[column] = None
    return work[columns]


def generate_psk_sheet(server_frames: list[pd.DataFrame], columns: list[str]) -> pd.DataFrame:
    size_cols = [col for col in ("Latency", "Avg. Size", "Read Size", "Write Size") if col in columns]
    iops_cols = [col for col in ("Total IOPS", "Read IOPS", "Write IOPS") if col in columns]
    iosize_parts: list[pd.DataFrame] = []
    iops_parts: list[pd.DataFrame] = []

    for df in server_frames:
        work = _as_datetime_frame(df)
        if "DateTime" not in work.columns:
            continue
        if size_cols:
            present = ["DateTime"] + [col for col in size_cols if col in work.columns]
            if len(present) > 1:
                iosize_parts.append(work[present])
        if iops_cols:
            present = ["DateTime"] + [col for col in iops_cols if col in work.columns]
            if len(present) > 1:
                iops_parts.append(work[present])

    if not iosize_parts and not iops_parts:
        return pd.DataFrame()

    iosize_merged = pd.DataFrame()
    iops_merged = pd.DataFrame()
    if iosize_parts:
        combined = pd.concat(iosize_parts, ignore_index=True)
        agg = {col: "mean" for col in size_cols if col in combined.columns}
        iosize_merged = combined.groupby("DateTime", as_index=False).agg(agg)
    if iops_parts:
        combined = pd.concat(iops_parts, ignore_index=True)
        agg = {col: "sum" for col in iops_cols if col in combined.columns}
        iops_merged = combined.groupby("DateTime", as_index=False).agg(agg)

    if not iosize_merged.empty and not iops_merged.empty:
        psk = pd.merge(iosize_merged, iops_merged, on="DateTime", how="outer")
    elif not iosize_merged.empty:
        psk = iosize_merged
    else:
        psk = iops_merged

    psk = psk.sort_values("DateTime")
    psk["DateTime"] = psk["DateTime"].astype(str)
    return _select_columns(psk, columns)


class CustomReportGenerator:
    def __init__(
        self,
        output_dir: str | os.PathLike[str],
        location_servers: dict[str, list[str]],
        server_data: dict[str, pd.DataFrame],
        *,
        sections: list[str],
        metrics: list[str],
        columns: dict[str, list[str]] | None = None,
        psk_columns: list[str] | None = None,
        static_dir: Path | None = None,
    ) -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.location_servers = location_servers
        self.server_data = server_data
        self.sections = set(sections)
        self.metrics = [metric_id for metric_id in metrics if metric_id in metric_by_id()]
        self.columns = columns or {}
        self.psk_columns = resolve_psk_columns(psk_columns)
        self.static_dir = static_dir
        self.tmp_columns = (
            combined_tmp_columns(self.metrics, self.columns)
            if self.metrics
            else list(TMP_COLUMNS)
        )
        self._used_sheet_names: set[str] = set()

    def _unique_sheet_name(self, name: str) -> str:
        base = _sheet_name(name)
        candidate = base
        suffix = 2
        while candidate in self._used_sheet_names:
            extra = f"_{suffix}"
            candidate = f"{base[: 31 - len(extra)]}{extra}"
            suffix += 1
        self._used_sheet_names.add(candidate)
        return candidate

    def _location_frames(self, servers: list[str]) -> list[pd.DataFrame]:
        frames = []
        for server in servers:
            df = self.server_data.get(server)
            if df is not None and not df.empty:
                frames.append(df)
        return frames

    def _write_metric_sheets(self, workbook: Workbook, location: str, server: str, df: pd.DataFrame) -> int:
        written = 0
        location_clean = _clean_location(location)
        catalog = metric_by_id()
        for metric_id in self.metrics:
            metric = catalog[metric_id]
            columns = resolve_metric_columns(metric_id, self.columns.get(metric_id))
            subset = _select_columns(df, columns)
            data_cols = [column for column in subset.columns if column != "DateTime"]
            if data_cols:
                subset = subset.dropna(how="all", subset=data_cols)
            present = [column for column in data_cols if subset[column].notna().any()]
            if subset.empty or not present:
                continue
            subset = subset[["DateTime"] + present]
            suffix = metric.get("sheet_suffix") or metric_id
            title = self._unique_sheet_name(f"{location_clean}_{server}_{suffix}")
            worksheet = workbook.create_sheet(title=title)
            for row in dataframe_to_rows(subset, index=False, header=True):
                worksheet.append(row)
            _style_metric_sheet(worksheet)
            written += 1
        return written

    def _write_tmp_sheet(self, workbook: Workbook, location: str, server: str, df: pd.DataFrame) -> None:
        tmp_df = build_server_tmp_table(location, server, df, columns=self.tmp_columns)
        title = self._unique_sheet_name(f"{_clean_location(location)}_{server}_TMP")
        worksheet = workbook.create_sheet(title=title)
        for row in dataframe_to_rows(tmp_df, index=False, header=False):
            worksheet.append(row)
        _style_tmp_sheet(worksheet)

    def _write_location_workbook(self, location: str, servers: list[str]) -> Path | None:
        frames = self._location_frames(servers)
        if not frames:
            return None

        wants_location = bool(self.sections & {"psk", "metric_sheets", "tmp_sheets"})
        if not wants_location:
            return None

        workbook = Workbook()
        workbook.remove(workbook.active)
        self._used_sheet_names = set()
        sheets = 0
        location_clean = _clean_location(location)

        if "psk" in self.sections:
            psk = generate_psk_sheet(frames, self.psk_columns)
            data_cols = [column for column in psk.columns if column != "DateTime"]
            if data_cols:
                psk = psk.dropna(how="all", subset=data_cols)
            if not psk.empty:
                title = self._unique_sheet_name("PSK")
                worksheet = workbook.create_sheet(title=title)
                for row in dataframe_to_rows(psk, index=False, header=True):
                    worksheet.append(row)
                _style_metric_sheet(worksheet)
                sheets += 1

        for server in servers:
            df = self.server_data.get(server)
            if df is None or df.empty:
                continue
            if "metric_sheets" in self.sections:
                sheets += self._write_metric_sheets(workbook, location, server, df)
            if "tmp_sheets" in self.sections:
                self._write_tmp_sheet(workbook, location, server, df)
                sheets += 1

        if sheets == 0:
            return None

        safe_location = re.sub(r"[^A-Za-z0-9_-]+", "_", location_clean) or "Location"
        path = self.output_dir / f"Custom_{safe_location}_Report.xlsx"
        workbook.save(path)
        return path

    def generate(self) -> dict[str, Any]:
        files: list[Path] = []
        for location, servers in self.location_servers.items():
            path = self._write_location_workbook(location, servers)
            if path is not None:
                files.append(path)

        if "all_tmps" in self.sections:
            hourly = HourlyReportGenerator(
                output_dir=self.output_dir,
                location_servers=self.location_servers,
                server_data=self.server_data,
                static_dir=self.static_dir,
                columns=self.tmp_columns,
                filename="Custom_All_TMPs.xlsx",
            )
            files.append(Path(hourly.generate()))

        if not files:
            raise ValueError("No custom report sheets could be generated from the selected options")

        if len(files) == 1:
            path = files[0]
            return {"output_file": str(path), "filename": path.name, "files": [path.name]}

        zip_path = self.output_dir / "Custom_Storage_Report.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in files:
                archive.write(path, path.name)
        return {
            "output_file": str(zip_path),
            "filename": zip_path.name,
            "files": [path.name for path in files],
        }
