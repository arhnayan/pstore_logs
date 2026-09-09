"""Generate All_Locations_Overprovisioning_Report.xlsx from cluster and volume space metrics."""

from __future__ import annotations

import os
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SUMMARY_HEADERS = [
    "Physical Total (TB)",
    "Logical Provisioned (TB)",
    "Logical Used (TB)",
    "Overprovisioning %",
    "Efficiency Ratio",
    "Thin Savings",
]

VOLUME_HEADERS = [
    "Server",
    "Volume Name",
    "Provisioned (TB)",
    "Logical Used (TB)",
    "Used %",
    "Type",
]


def _sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "", name).strip() or "Sheet"
    return cleaned[:31]


class OverprovisionReportGenerator:
    RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    BLACK_FILL = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    GREY_FILL = PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")
    YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    LIGHT_RED_FILL = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
    WHITE_FONT = Font(color="FFFFFF", bold=True)
    BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def __init__(
        self,
        output_dir: str,
        location_servers: dict[str, list[str]],
        server_space: dict[str, dict[str, float]] | None = None,
        volume_rows: dict[str, list[dict[str, Any]]] | None = None,
    ) -> None:
        self.output_dir = output_dir
        self.location_servers = location_servers
        self.server_space = server_space or {}
        self.volume_rows = volume_rows or {}
        os.makedirs(output_dir, exist_ok=True)
        self._used_sheet_names: set[str] = set()

    def _unique_sheet_name(self, name: str) -> str:
        base = _sheet_name(name)
        candidate = base
        suffix = 2
        while candidate in self._used_sheet_names:
            extra = f"_{suffix}"
            candidate = f"{base[: 31 - len(extra)]}{extra}"
            suffix += 1
        self._used_sheet_names.add(candidate)
        return candidate

    def _style_header_cell(self, cell, fill) -> None:
        cell.fill = fill
        cell.font = self.WHITE_FONT if fill in (self.RED_FILL, self.BLACK_FILL) else Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = self.BORDER

    def _write_grey_headers(self, ws, row: int, headers: list[str]) -> None:
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.value = header
            self._style_header_cell(cell, self.GREY_FILL)

    def _format_ratio(self, value: Any) -> str | None:
        if value is None or value == "":
            return None
        try:
            return f"{float(value):.1f}:1"
        except (TypeError, ValueError):
            return None

    def _apply_overprovision_fill(self, cell, pct: float | None) -> None:
        if pct is None:
            return
        if pct > 100:
            cell.fill = self.LIGHT_RED_FILL
        elif pct > 80:
            cell.fill = self.YELLOW_FILL

    def add_location_summary_sheet(
        self,
        wb: Workbook,
        location: str,
        servers: list[str],
    ) -> None:
        ws = wb.create_sheet(title=self._unique_sheet_name(f"{location} Report"))
        for col in range(1, 7):
            ws.column_dimensions[get_column_letter(col)].width = 22

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        header = ws.cell(row=1, column=1)
        header.value = f"{location} Site  {', '.join(servers)}"
        header.fill = self.RED_FILL
        header.font = self.WHITE_FONT
        header.alignment = Alignment(horizontal="center", vertical="center")

        current_row = 3
        for server in servers:
            current_row = self._create_server_section(ws, server, current_row)

    def _create_server_section(self, ws, server_name: str, start_row: int) -> int:
        current_row = start_row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        name_cell = ws.cell(row=current_row, column=1)
        name_cell.value = server_name
        name_cell.fill = self.BLACK_FILL
        name_cell.font = self.WHITE_FONT
        name_cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        sub = ws.cell(row=current_row, column=1)
        sub.value = "OVERPROVISIONING"
        self._style_header_cell(sub, self.GREY_FILL)
        current_row += 1

        self._write_grey_headers(ws, current_row, SUMMARY_HEADERS)
        current_row += 1

        space = self.server_space.get(server_name) or self.server_space.get(server_name.upper(), {})
        physical = space.get("Physical_Total_TB")
        provisioned = space.get("Logical_Provisioned_TB")
        used = space.get("Logical_Used_TB")
        over_pct = space.get("Overprovisioning_Pct")
        values = [
            physical,
            provisioned,
            used,
            over_pct,
            self._format_ratio(space.get("Efficiency_Ratio")),
            self._format_ratio(space.get("Thin_Savings")),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = value
            cell.border = self.BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx <= 3 and isinstance(value, (int, float)):
                cell.number_format = "0.00"
            if col_idx == 4 and isinstance(value, (int, float)):
                cell.number_format = "0.0"
                self._apply_overprovision_fill(cell, float(value))
        return current_row + 2

    def add_location_volume_sheet(
        self,
        wb: Workbook,
        location: str,
        volume_rows: list[dict[str, Any]],
    ) -> None:
        ws = wb.create_sheet(title=self._unique_sheet_name(f"{location} Volumes"))
        widths = [22, 36, 20, 20, 14, 14]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        header = ws.cell(row=1, column=1)
        header.value = f"{location} Volumes"
        header.fill = self.RED_FILL
        header.font = self.WHITE_FONT
        header.alignment = Alignment(horizontal="center", vertical="center")

        self._write_grey_headers(ws, 2, VOLUME_HEADERS)
        rows = sorted(
            volume_rows,
            key=lambda row: (str(row.get("server") or ""), str(row.get("volume_name") or "")),
        )
        for idx, row in enumerate(rows):
            excel_row = idx + 3
            values = [
                row.get("server"),
                row.get("volume_name"),
                row.get("provisioned_tb"),
                row.get("logical_used_tb"),
                row.get("used_pct"),
                row.get("type"),
            ]
            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=excel_row, column=col_idx)
                cell.value = value
                cell.border = self.BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_idx in (3, 4) and isinstance(value, (int, float)):
                    cell.number_format = "0.00"
                if col_idx == 5 and isinstance(value, (int, float)):
                    cell.number_format = "0.0"

    def generate(self) -> str:
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._used_sheet_names.clear()

        for location, servers in self.location_servers.items():
            self.add_location_summary_sheet(workbook, location, servers)
            self.add_location_volume_sheet(
                workbook,
                location,
                self.volume_rows.get(location, []),
            )

        output_file = os.path.join(self.output_dir, "All_Locations_Overprovisioning_Report.xlsx")
        workbook.save(output_file)
        return output_file
